"""
Multi-Format Document Loader with Text Cleaning.
Supports PDF (streaming up to ~1000 pages), DOCX, TXT, MD, CSV, PPTX, and XLSX.
"""
from __future__ import annotations
import re
import csv
import logging
from pathlib import Path
from typing import List, Dict, Any

logger = logging.getLogger(__name__)


def clean_text(text: str) -> str:
    """
    Clean unnecessary whitespace and malformed characters while
    preserving meaningful document and paragraph structure.
    """
    if not text:
        return ""

    # Remove null bytes and control characters (except newline, tab, carriage return)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)

    # Rejoin words broken across line breaks with hyphens: "inter-\nnational" -> "international"
    text = re.sub(r"(\b\w+)-\n(\w+\b)", r"\1\2", text)

    # Normalize carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Normalize horizontal whitespace (tabs and multiple spaces)
    lines = []
    for line in text.split("\n"):
        line_clean = re.sub(r"[ \t]+", " ", line).strip()
        lines.append(line_clean)

    # Collapse more than 2 consecutive newlines into 2 (preserve paragraph breaks)
    cleaned = "\n".join(lines)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)

    return cleaned.strip()


def load_document(file_path: str | Path) -> List[Dict[str, Any]]:
    """
    Extract structured pages/records from any supported document format.

    Returns a list of dicts:
        [
            {
                "page": int,
                "text": str,
                "metadata": {
                    "source": str,
                    "page_number": int,
                    "total_pages": int,
                    "file_type": str,
                    "section": str (optional)
                }
            }
        ]
    """
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _load_pdf(path)
    elif suffix in (".txt", ".md", ".markdown"):
        return _load_text(path)
    elif suffix == ".docx":
        return _load_docx(path)
    elif suffix == ".pptx":
        return _load_pptx(path)
    elif suffix == ".csv":
        return _load_csv(path)
    elif suffix in (".xlsx", ".xls"):
        return _load_xlsx(path)
    else:
        raise ValueError(
            f"Unsupported file format: '{suffix}'. Supported formats: PDF, DOCX, TXT, MD, CSV, PPTX, XLSX"
        )


def _load_pdf(path: Path) -> List[Dict[str, Any]]:
    """
    Extract PDF text efficiently using PyMuPDF (fitz) or pypdf fallback.
    Streams page-by-page to handle large documents (approaching 1000 pages)
    without memory bloat.
    """
    pages: List[Dict[str, Any]] = []

    # Attempt 1: PyMuPDF (Fastest, low memory footprint)
    try:
        import pymupdf  # PyMuPDF
        doc = pymupdf.open(str(path))
        total_pages = len(doc)
        logger.info(f"Extracting {total_pages} pages from PDF via PyMuPDF: {path.name}")

        for i in range(total_pages):
            page = doc[i]
            raw_text = page.get_text() or ""
            cleaned = clean_text(raw_text)
            if cleaned:
                pages.append({
                    "page": i + 1,
                    "text": cleaned,
                    "metadata": {
                        "source": path.name,
                        "page_number": i + 1,
                        "total_pages": total_pages,
                        "file_type": "pdf",
                    },
                })
        doc.close()
        if pages:
            return pages
    except ImportError:
        logger.info("pymupdf not installed; falling back to pypdf / pdfplumber.")
    except Exception as e:
        logger.warning(f"PyMuPDF extraction failed ({e}); falling back to pypdf.")

    # Attempt 2: pypdf fallback
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        total_pages = len(reader.pages)
        logger.info(f"Extracting {total_pages} pages from PDF via pypdf: {path.name}")

        for i, page in enumerate(reader.pages):
            raw_text = page.extract_text() or ""
            cleaned = clean_text(raw_text)
            if cleaned:
                pages.append({
                    "page": i + 1,
                    "text": cleaned,
                    "metadata": {
                        "source": path.name,
                        "page_number": i + 1,
                        "total_pages": total_pages,
                        "file_type": "pdf",
                    },
                })
        if pages:
            return pages
    except Exception as e:
        logger.warning(f"pypdf extraction failed ({e}); falling back to pdfplumber.")

    # Attempt 3: pdfplumber fallback
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            total_pages = len(pdf.pages)
            for i, page in enumerate(pdf.pages):
                raw_text = page.extract_text() or ""
                cleaned = clean_text(raw_text)
                if cleaned:
                    pages.append({
                        "page": i + 1,
                        "text": cleaned,
                        "metadata": {
                            "source": path.name,
                            "page_number": i + 1,
                            "total_pages": total_pages,
                            "file_type": "pdf",
                        },
                    })
    except Exception as e:
        logger.error(f"All PDF extractors failed for {path.name}: {e}")
        raise ValueError(f"Could not read PDF file {path.name}: {e}")

    return pages


def _load_text(path: Path) -> List[Dict[str, Any]]:
    """Extract text from TXT or Markdown file, segmenting large files into pseudo-pages."""
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="latin-1", errors="replace")

    cleaned = clean_text(text)
    if not cleaned:
        return []

    # Segment into pseudo-pages of ~3000 characters for consistent indexing
    page_size = 3000
    pages: List[Dict[str, Any]] = []
    total_len = len(cleaned)
    estimated_pages = max(1, (total_len + page_size - 1) // page_size)

    for i in range(0, total_len, page_size):
        chunk = cleaned[i : i + page_size]
        if chunk.strip():
            page_num = len(pages) + 1
            pages.append({
                "page": page_num,
                "text": chunk,
                "metadata": {
                    "source": path.name,
                    "page_number": page_num,
                    "total_pages": estimated_pages,
                    "file_type": path.suffix.lower().lstrip("."),
                },
            })

    return pages


def _load_docx(path: Path) -> List[Dict[str, Any]]:
    """Extract paragraphs and tables from DOCX files."""
    from docx import Document

    doc = Document(str(path))
    content_blocks = []

    # Extract paragraphs
    for p in doc.paragraphs:
        txt = clean_text(p.text)
        if txt:
            content_blocks.append(txt)

    # Extract tables
    for t in doc.tables:
        table_rows = []
        for row in t.rows:
            row_text = " | ".join(clean_text(cell.text) for cell in row.cells if cell.text.strip())
            if row_text:
                table_rows.append(row_text)
        if table_rows:
            content_blocks.append("\n".join(table_rows))

    full_text = "\n\n".join(content_blocks)
    if not full_text.strip():
        return []

    # Segment into pseudo-pages
    page_size = 3000
    pages: List[Dict[str, Any]] = []
    total_len = len(full_text)
    estimated_pages = max(1, (total_len + page_size - 1) // page_size)

    for i in range(0, total_len, page_size):
        chunk = full_text[i : i + page_size]
        if chunk.strip():
            page_num = len(pages) + 1
            pages.append({
                "page": page_num,
                "text": chunk,
                "metadata": {
                    "source": path.name,
                    "page_number": page_num,
                    "total_pages": estimated_pages,
                    "file_type": "docx",
                },
            })

    return pages


def _load_pptx(path: Path) -> List[Dict[str, Any]]:
    """Extract slides, titles, shapes, and notes from PPTX files."""
    from pptx import Presentation

    prs = Presentation(str(path))
    pages: List[Dict[str, Any]] = []
    total_slides = len(prs.slides)

    for idx, slide in enumerate(prs.slides):
        slide_texts = []
        slide_title = ""

        # Extract shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = clean_text(shape.text)
                if txt:
                    if shape == slide.shapes[0] and not slide_title:
                        slide_title = txt
                    slide_texts.append(txt)

        # Extract notes if any
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes_text:
                slide_texts.append(f"[Slide Notes: {notes_text}]")

        combined = "\n".join(slide_texts).strip()
        if combined:
            pages.append({
                "page": idx + 1,
                "text": combined,
                "metadata": {
                    "source": path.name,
                    "page_number": idx + 1,
                    "total_pages": total_slides,
                    "file_type": "pptx",
                    "section": slide_title[:80] if slide_title else f"Slide {idx + 1}",
                },
            })

    return pages


def _load_csv(path: Path) -> List[Dict[str, Any]]:
    """Extract CSV records grouping rows into readable chunks with headers."""
    import pandas as pd

    try:
        df = pd.read_csv(str(path))
    except Exception:
        df = pd.read_csv(str(path), encoding="latin-1")

    pages: List[Dict[str, Any]] = []
    if df.empty:
        return pages

    total_rows = len(df)
    batch_size = 20  # 20 rows per chunk/record-page
    total_batches = (total_rows + batch_size - 1) // batch_size

    for batch_idx, start_idx in enumerate(range(0, total_rows, batch_size)):
        chunk_df = df.iloc[start_idx : start_idx + batch_size]
        text_records = []
        for _, row in chunk_df.iterrows():
            record_str = ", ".join(f"{col}: {val}" for col, val in row.items() if pd.notna(val))
            text_records.append(record_str)

        combined = "\n".join(text_records)
        pages.append({
            "page": batch_idx + 1,
            "text": clean_text(combined),
            "metadata": {
                "source": path.name,
                "page_number": batch_idx + 1,
                "total_pages": total_batches,
                "file_type": "csv",
                "section": f"Records {start_idx + 1} - {min(start_idx + batch_size, total_rows)}",
            },
        })

    return pages


def _load_xlsx(path: Path) -> List[Dict[str, Any]]:
    """Extract all sheets and rows from XLSX / XLS files."""
    import pandas as pd

    excel_file = pd.ExcelFile(str(path))
    pages: List[Dict[str, Any]] = []
    page_counter = 1

    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(excel_file, sheet_name=sheet_name)
        if df.empty:
            continue

        total_rows = len(df)
        batch_size = 25
        for start_idx in range(0, total_rows, batch_size):
            chunk_df = df.iloc[start_idx : start_idx + batch_size]
            text_records = []
            for _, row in chunk_df.iterrows():
                record_str = ", ".join(f"{col}: {val}" for col, val in row.items() if pd.notna(val))
                text_records.append(record_str)

            combined = f"Sheet: {sheet_name}\n" + "\n".join(text_records)
            pages.append({
                "page": page_counter,
                "text": clean_text(combined),
                "metadata": {
                    "source": path.name,
                    "page_number": page_counter,
                    "file_type": "xlsx",
                    "section": f"Sheet: {sheet_name} (Rows {start_idx + 1} - {min(start_idx + batch_size, total_rows)})",
                },
            })
            page_counter += 1

    # Update total pages in metadata
    for p in pages:
        p["metadata"]["total_pages"] = len(pages)

    return pages
