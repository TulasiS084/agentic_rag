"""
Utility script to generate sample test documents (TXT, DOCX, CSV, PDF) for testing.
"""
import os
from pathlib import Path

def generate_samples():
    sample_dir = Path(__file__).parent.parent / "sample_docs"
    sample_dir.mkdir(exist_ok=True)

    # 1. Text Document
    txt_content = """Artificial Intelligence and Machine Learning: A Comprehensive Guide

Chapter 1: Introduction to Artificial Intelligence
Artificial Intelligence (AI) is a multidisciplinary field of computer science dedicated to developing computer systems capable of performing tasks that ordinarily require human intelligence. Such cognitive abilities encompass visual pattern perception, speech interpretation, automated decision-making, and natural language translation.

Chapter 2: Machine Learning Foundations
Machine Learning (ML) is the computational backbone of modern AI. ML algorithms discover statistical relationships and patterns directly from training data rather than relying on deterministic, hard-coded rule sets.
Key paradigms include:
- Supervised Learning: Learning mappings from input features to known target labels (e.g. classification, regression).
- Unsupervised Learning: Extracting inherent cluster structures or low-dimensional representations from unlabelled data.
- Reinforcement Learning: Optimizing agent policies through trial-and-error environmental interactions driven by scalar reward signals.

Chapter 3: Deep Neural Networks and Architectures
Deep Learning models hierarchically compose representation layers using artificial neural network topologies:
- Convolutional Neural Networks (CNNs): Excel at spatial feature hierarchies, dominating computer vision, medical imaging, and object recognition.
- Recurrent Neural Networks (RNNs) and LSTMs: Process sequential dependencies across temporal sequences.
- Transformer Architectures: Leverage self-attention mechanisms to parallelize contextual token representations, serving as the foundation of modern Large Language Models (LLMs) such as GPT, Claude, and Llama.
- Generative Adversarial Networks (GANs): Pair generator and discriminator sub-networks in a minimax optimization game to synthesize hyper-realistic synthetic media.

Chapter 4: Agentic RAG Systems
Retrieval-Augmented Generation (RAG) augments generative models with external factual databases, dramatically reducing hallucinations.
Agentic RAG elevates this paradigm by integrating autonomous query routing, hybrid dense-sparse retrieval (combining dense embeddings with BM25), reciprocal rank fusion (RRF), and cross-encoder reranking to ensure that retrieved context is maximally relevant, precise, and verifiable with exact source citations.
"""
    (sample_dir / "AI_and_Machine_Learning.txt").write_text(txt_content.strip(), encoding="utf-8")
    print(f"[OK] Generated {sample_dir / 'AI_and_Machine_Learning.txt'}")

    # 2. DOCX Document
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("Deep Learning Architectural Guide", level=0)
        doc.add_heading("Section 1: Foundations of Artificial Neural Networks", level=1)
        doc.add_paragraph(
            "An Artificial Neural Network (ANN) consists of input, hidden, and output layers of interconnected neurons. "
            "Each connection possesses a tunable synaptic weight adjusted via backpropagation and gradient descent."
        )
        doc.add_heading("Section 2: The Attention Mechanism", level=1)
        doc.add_paragraph(
            "The Transformer architecture introduced self-attention: Attention(Q, K, V) = softmax(Q K^T / sqrt(d_k)) * V. "
            "This mathematical formulation enables models to weight every input token relative to all other tokens in a single parallel operation."
        )
        doc.save(str(sample_dir / "Deep_Learning_Guide.docx"))
        print(f"[OK] Generated {sample_dir / 'Deep_Learning_Guide.docx'}")
    except Exception as e:
        print(f"Failed to generate docx: {e}")

    # 3. CSV Document
    try:
        import pandas as pd
        csv_data = {
            "Department": ["Research & Development", "Engineering", "Marketing", "Human Resources", "Operations", "Finance"],
            "Budget_Millions": [14.5, 28.2, 8.7, 3.4, 11.2, 5.8],
            "Headcount": [45, 120, 35, 15, 60, 22],
            "Quarterly_Growth_Percent": [18.5, 24.1, 9.4, 4.2, 12.0, 7.3],
            "Key_Initiative": [
                "Agentic RAG and Multimodal LLMs",
                "Distributed Vector Databases and Infrastructure",
                "Global Developer Community Outreach",
                "Technical Talent Retention and Upskilling",
                "Cloud Compute Cost Optimization",
                "Automated Financial Auditing Pipelines",
            ]
        }
        df = pd.DataFrame(csv_data)
        df.to_csv(sample_dir / "Quarterly_Department_Metrics.csv", index=False)
        print(f"[OK] Generated {sample_dir / 'Quarterly_Department_Metrics.csv'}")
    except Exception as e:
        print(f"Failed to generate CSV: {e}")

    # 4. PDF Document using PyMuPDF
    try:
        import pymupdf
        pdf_doc = pymupdf.open()
        
        # Page 1
        page1 = pdf_doc.new_page()
        page1.insert_text(
            (50, 70),
            "Executive Report: Advanced Retrieval-Augmented Generation\n\n"
            "1. Executive Summary\n"
            "This report analyzes the performance benchmarks of Agentic RAG compared to standard vector retrieval.\n"
            "Our findings indicate that hybrid dense-sparse retrieval with cross-encoder reranking yields a 38% improvement\n"
            "in factual accuracy and eliminates 92% of semantic hallucination risks.\n\n"
            "2. Hybrid Retrieval Performance\n"
            "Dense vector embeddings (all-MiniLM-L6-v2) capture semantic meaning and conceptual similarity.\n"
            "Sparse lexical retrieval (BM25) guarantees exact keyword and numerical match precision.\n"
            "Reciprocal Rank Fusion (RRF) harmonizes both rankings using a constant k=60.",
            fontsize=12,
        )

        # Page 2
        page2 = pdf_doc.new_page()
        page2.insert_text(
            (50, 70),
            "3. Reranking and Context Assembly\n\n"
            "The cross-encoder/ms-marco-MiniLM-L-6-v2 evaluates query-document pairs simultaneously.\n"
            "This resolves subtle context nuances that bi-encoders miss.\n"
            "Only the top 3-5 reranked chunks are passed to the LLM generator to maintain concise prompts.\n\n"
            "4. Conclusion\n"
            "The resulting system delivers production-grade enterprise research capabilities with full citation tracking.",
            fontsize=12,
        )

        pdf_path = sample_dir / "Executive_RAG_Report.pdf"
        pdf_doc.save(str(pdf_path))
        pdf_doc.close()
        print(f"[OK] Generated {pdf_path}")
    except Exception as e:
        print(f"Failed to generate PDF: {e}")

if __name__ == "__main__":
    generate_samples()
